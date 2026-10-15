"""체크리스트를 표 형식 워드 문서(.docx)로 만든다.

열 구성: 선택 | 번호 | 단계 | 검토 분야 | 체크리스트 | 확인 필요사항 | 필요 산출물 | 담당자 | 소속
머리행은 남색 배경 + 흰 글씨, 본문은 한 줄씩 옅은 회청색 줄무늬.
문서 순서: 제목·프로젝트 개요 → 프로젝트 일정(예상, 간트 표) → 카테고리별 체크리스트
          → <첨부> 규정 근거 → <첨부> 필요 산출물 서식.
"""
import io
from datetime import date, timedelta

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from timeline import BASE_WEIGHTS

FONT = "맑은 고딕"
HEADER_FILL = "1F3864"
STRIPE_FILL = "EEF3F8"
BORDER_COLOR = "BFC8D6"
NUMBER_COLOR = RGBColor(0x1F, 0x38, 0x64)

# (머리글, 폭 cm, 가운데 정렬 여부) — 가로 A4 본문 폭(약 26.7cm)에 맞춤
COLUMNS = [
    ("선택", 1.2, True),
    ("번호", 1.5, True),
    ("단계", 1.7, True),
    ("검토 분야", 2.4, True),
    ("체크리스트", 6.9, False),
    ("확인 필요사항", 4.2, False),
    ("필요 산출물", 4.0, False),
    ("담당자", 2.0, True),
    ("소속", 2.8, False),
]
PHASE_ORDER = ["기획", "설계", "개발", "테스트", "오픈"]


def _shade(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _set_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), BORDER_COLOR)
        borders.append(el)
    tbl_pr.append(borders)


def _row_flag(row, tag: str):
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement(f"w:{tag}")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def _write(cell, text: str, *, center: bool, bold=False, color=None, size=9.5, font=FONT,
           keep_with_next=False):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    para = cell.paragraphs[0]
    para.paragraph_format.keep_with_next = keep_with_next
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    para.paragraph_format.space_before = Pt(3)
    para.paragraph_format.space_after = Pt(3)
    run = para.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = font
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), font)
    if color is not None:
        run.font.color.rgb = color


def _setup_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Cm(1.5))

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    return doc


def _heading(doc, text: str, size: float, space_after: float):
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(space_after)
    para.paragraph_format.keep_with_next = True
    run = para.add_run(text)
    run.bold = True
    run.font.size = Pt(size)
    run.font.color.rgb = NUMBER_COLOR
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)


def _add_table(doc, items: list[dict]):
    table = doc.add_table(rows=1, cols=len(COLUMNS))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _set_borders(table)

    header = table.rows[0]
    _row_flag(header, "tblHeader")  # 페이지가 넘어가면 머리행 반복
    for cell, (label, width, _) in zip(header.cells, COLUMNS):
        cell.width = Cm(width)
        _shade(cell, HEADER_FILL)
        # 머리행이 페이지 끝에 홀로 남지 않도록 첫 데이터 행과 붙여 둔다.
        _write(cell, label, center=True, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), keep_with_next=True)

    for idx, item in enumerate(items):
        owner = item.get("owner") or {}
        values = [
            "☑" if item.get("done") else "☐",
            item.get("no", ""),
            item.get("phase") or "-",
            # 리스크 항목은 검토 분야 아래 위험도, 체크리스트 아래 AI 판단 근거를 붙인다.
            (item.get("area") or "-") + (f"\n위험도 {item['level']}" if item.get("level") else ""),
            item.get("text", "") + (f"\n└ 판단 근거: {item['reason']}" if item.get("reason") else ""),
            item.get("check") or "-",
            "\n".join(f"[{d['code']}] {d['name']}" for d in item.get("deliverables") or []) or "-",
            # 담당자(이름/직급)·소속(부서/팀)은 좁은 칸에서 어색하게 끊기지 않도록 두 줄로 쓴다.
            "\n".join(p for p in (owner.get("name"), owner.get("title")) if p) if owner else "미지정",
            "\n".join(p for p in (owner.get("dept"), owner.get("team")) if p) or "-",
        ]
        row = table.add_row()
        _row_flag(row, "cantSplit")  # 한 행이 두 페이지로 쪼개지지 않게
        for col, (cell, value, (_, width, center)) in enumerate(zip(row.cells, values, COLUMNS)):
            cell.width = Cm(width)
            if idx % 2 == 1:
                _shade(cell, STRIPE_FILL)
            if col == 0:
                _write(cell, value, center=True, size=13, font="Segoe UI Symbol")
            elif col == 1:
                _write(cell, value, center=True, bold=True, color=NUMBER_COLOR)
            else:
                _write(cell, value, center=center)


# "<첨부> 규정 근거" 표: 번호 | 검토 분야 | 체크리스트 | 근거 규정·조항
BASIS_COLUMNS = [
    ("번호", 1.8, True),
    ("검토 분야", 3.0, True),
    ("체크리스트", 10.4, False),
    ("근거 규정·조항", 11.5, False),
]
BASIS_NOTE = ("※ 근거 조항은 항목 내용을 기준으로 자동 매칭한 참고 자료입니다. "
              "최종 적용 여부와 조항 번호는 준법감시부 확인 후 확정하세요.")


def _add_basis_appendix(doc, categories: list[dict]):
    """체크리스트 표 뒤에 새 페이지로 항목별 근거 규정을 붙인다."""
    rows = [item for c in categories for item in c["items"] if item.get("basis")]
    if not rows:
        return
    doc.add_page_break()
    _heading(doc, "<첨부> 규정 근거", size=14, space_after=8)

    table = doc.add_table(rows=1, cols=len(BASIS_COLUMNS))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _set_borders(table)
    header = table.rows[0]
    _row_flag(header, "tblHeader")
    for cell, (label, width, _) in zip(header.cells, BASIS_COLUMNS):
        cell.width = Cm(width)
        _shade(cell, HEADER_FILL)
        _write(cell, label, center=True, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), keep_with_next=True)

    for idx, item in enumerate(rows):
        values = [item.get("no", ""), item.get("area") or "-", item.get("text", ""),
                  "\n".join(f"· {b}" for b in item["basis"])]
        row = table.add_row()
        _row_flag(row, "cantSplit")
        for col, (cell, value, (_, width, center)) in enumerate(zip(row.cells, values, BASIS_COLUMNS)):
            cell.width = Cm(width)
            if idx % 2 == 1:
                _shade(cell, STRIPE_FILL)
            _write(cell, value, center=center, bold=col == 0, color=NUMBER_COLOR if col == 0 else None)

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(6)
    run = note.add_run(BASIS_NOTE)
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)


# "프로젝트 개요" 표: 영역 | 항목 | 내용 | 확인 방식
PROFILE_COLUMNS = [
    ("영역", 2.6, True),
    ("항목", 5.0, True),
    ("내용", 12.6, False),
    ("확인 방식", 6.5, False),
]
SOURCE_LABEL = {"inferred": "질문에서 파악", "answered": "사용자 응답"}


def _add_profile_section(doc, project: dict):
    """체크리스트 앞에 프로파일링으로 수집한 프로젝트 기본 정보를 표로 넣는다."""
    _heading(doc, "프로젝트 개요", size=11.5, space_after=4)
    table = doc.add_table(rows=1, cols=len(PROFILE_COLUMNS))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _set_borders(table)
    header = table.rows[0]
    _row_flag(header, "tblHeader")
    for cell, (label, width, _) in zip(header.cells, PROFILE_COLUMNS):
        cell.width = Cm(width)
        _shade(cell, HEADER_FILL)
        _write(cell, label, center=True, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), keep_with_next=True)

    for idx, item in enumerate(project.get("items", [])):
        source = SOURCE_LABEL.get(item.get("source"), "사용자 응답")
        if item.get("value") == "미확인":
            source = "확인 필요 (프로파일링 시 미응답)"
        elif item.get("source") == "inferred" and item.get("reason"):
            source += f" ({item['reason']})"
        row = table.add_row()
        _row_flag(row, "cantSplit")
        values = [item.get("area", ""), item.get("label", ""), item.get("value", ""), source]
        for cell, value, (_, width, center) in zip(row.cells, values, PROFILE_COLUMNS):
            cell.width = Cm(width)
            if idx % 2 == 1:
                _shade(cell, STRIPE_FILL)
            _write(cell, value, center=center)
    doc.add_paragraph()


# "<첨부> 필요 산출물 서식" 표: 서식번호 | 서식명 | 작성 단계 | 관련 체크리스트
FORM_COLUMNS = [
    ("서식번호", 2.6, True),
    ("서식명", 9.0, False),
    ("작성 단계", 3.0, True),
    ("관련 체크리스트", 12.1, False),
]
FORM_NOTE = ("※ 서식과 작성 단계는 「IT 프로젝트 산출물 관리지침(YH-IT-007)」 별표1·별지 서식 기준입니다.")


def _add_forms_appendix(doc, categories: list[dict]):
    """체크리스트에 걸린 산출물 서식을 작성 단계 순으로 모아 서식별 관련 항목 번호를 보여준다."""
    forms: dict[str, dict] = {}
    for cat in categories:
        for item in cat["items"]:
            for d in item.get("deliverables") or []:
                f = forms.setdefault(d["code"], {"name": d["name"], "phase": d.get("phase", ""), "nos": []})
                f["nos"].append(item.get("no", ""))
    if not forms:
        return
    doc.add_page_break()
    _heading(doc, "<첨부> 필요 산출물 서식", size=14, space_after=8)

    rows = sorted(forms.items(), key=lambda kv: (
        PHASE_ORDER.index(kv[1]["phase"]) if kv[1]["phase"] in PHASE_ORDER else 99, kv[0]))
    table = doc.add_table(rows=1, cols=len(FORM_COLUMNS))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _set_borders(table)
    header = table.rows[0]
    _row_flag(header, "tblHeader")
    for cell, (label, width, _) in zip(header.cells, FORM_COLUMNS):
        cell.width = Cm(width)
        _shade(cell, HEADER_FILL)
        _write(cell, label, center=True, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), keep_with_next=True)
    for idx, (code, f) in enumerate(rows):
        values = [code, f["name"], f["phase"] or "-", ", ".join(n for n in f["nos"] if n)]
        row = table.add_row()
        _row_flag(row, "cantSplit")
        for col, (cell, value, (_, width, center)) in enumerate(zip(row.cells, values, FORM_COLUMNS)):
            cell.width = Cm(width)
            if idx % 2 == 1:
                _shade(cell, STRIPE_FILL)
            _write(cell, value, center=center, bold=col == 0, color=NUMBER_COLOR if col == 0 else None)

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(6)
    run = note.add_run(FORM_NOTE)
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)


# "프로젝트 일정(예상)" 간트 표. 막대 색은 화면(App.css --phase-*)과 같은 단계별 색.
PHASE_FILL = {"기획": "2A78D6", "설계": "EB6834", "개발": "1BAF7A", "테스트": "EDA100", "오픈": "E87BA4"}
GANTT_FIXED = [("단계", 1.5), ("기간", 4.0), ("일수", 1.5)]
GANTT_TRACK_WIDTH = 13.2   # 기간 칸(주·월·분기)들이 나눠 쓰는 폭
GANTT_FORMS_WIDTH = 6.5    # 주요 산출물 칸


def _periods(start: date, end: date) -> list[tuple[str, date, date]]:
    """간트 표의 기간 칸: 14주 이하면 주, 18개월 이하면 월, 그보다 길면 분기."""
    if (end - start).days + 1 <= 98:
        out, cur, n = [], start, 1
        while cur <= end:
            out.append((f"{n}주", cur, min(cur + timedelta(days=6), end)))
            cur += timedelta(days=7)
            n += 1
        return out
    months = []
    y, m = start.year, start.month
    while (y, m) <= (end.year, end.month):
        first = date(y, m, 1)
        nxt = date(y + (m == 12), m % 12 + 1, 1)
        months.append((f"{y % 100}.{m:02d}", max(first, start), min(nxt - timedelta(days=1), end)))
        y, m = nxt.year, nxt.month
    if len(months) <= 18:
        return months
    quarters: dict[str, list] = {}
    for _, s, e in months:
        key = f"{s.year % 100} Q{(s.month - 1) // 3 + 1}"
        q = quarters.setdefault(key, [s, e])
        q[1] = e
    return [(k, s, e) for k, (s, e) in quarters.items()]


def _add_timeline_section(doc, timeline: dict, categories: list[dict]):
    """단계별 예상 일정을 간트 표로 그린다. 막대 = 단계 기간과 겹치는 기간 칸의 음영."""
    _heading(doc, "프로젝트 일정(예상)", size=11.5, space_after=4)
    periods = _periods(date.fromisoformat(timeline["start"]), date.fromisoformat(timeline["end"]))
    track_w = GANTT_TRACK_WIDTH / len(periods)

    # 단계별 주요 산출물: 체크리스트에 걸린 서식 중 작성 단계가 그 단계인 것
    forms_by_phase: dict[str, list[str]] = {}
    for cat in categories:
        for item in cat["items"]:
            for d in item.get("deliverables") or []:
                codes = forms_by_phase.setdefault(d.get("phase", ""), [])
                if d["code"] not in codes:
                    codes.append(d["code"])

    widths = [w for _, w in GANTT_FIXED] + [track_w] * len(periods) + [GANTT_FORMS_WIDTH]
    labels = [l for l, _ in GANTT_FIXED] + [p[0] for p in periods] + ["주요 산출물(작성 서식)"]
    table = doc.add_table(rows=1, cols=len(widths))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _set_borders(table)
    header = table.rows[0]
    _row_flag(header, "tblHeader")
    for i, (cell, label, width) in enumerate(zip(header.cells, labels, widths)):
        cell.width = Cm(width)
        _shade(cell, HEADER_FILL)
        is_period = len(GANTT_FIXED) <= i < len(GANTT_FIXED) + len(periods)
        _write(cell, label, center=True, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF),
               size=7.5 if is_period else 9.5, keep_with_next=True)

    for ph in timeline["phases"]:
        s, e = date.fromisoformat(ph["start"]), date.fromisoformat(ph["end"])
        row = table.add_row()
        _row_flag(row, "cantSplit")
        cells = row.cells
        fixed = [ph["phase"], f"{ph['start']} ~ {ph['end']}", f"{ph['days']}일"]
        for cell, value, width in zip(cells, fixed, widths):
            cell.width = Cm(width)
            _write(cell, value, center=True, bold=value == ph["phase"])
        for j, (_, ps, pe) in enumerate(periods):
            cell = cells[len(GANTT_FIXED) + j]
            cell.width = Cm(track_w)
            if s <= pe and e >= ps:  # 단계 기간과 겹치는 칸에 막대 음영
                _shade(cell, PHASE_FILL[ph["phase"]])
        forms = forms_by_phase.get(ph["phase"], [])
        cells[-1].width = Cm(GANTT_FORMS_WIDTH)
        _write(cells[-1], ", ".join(forms) or "-", center=False, size=8.5)

    base = " · ".join(f"{p} {w}%" for p, w in BASE_WEIGHTS.items())
    lines = [f"※ 총 {timeline['total_days']}일(달력일). 기본 비중 {base}에서 프로파일에 따라 조정했습니다."]
    lines += [f"   - {n}" for n in timeline["notes"]]
    for text in lines:
        para = doc.add_paragraph()
        para.paragraph_format.space_after = Pt(0)
        run = para.add_run(text)
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    doc.add_paragraph()


def build_checklist_docx(categories: list[dict], title: str = "IT 프로젝트 검토 체크리스트",
                         project: dict | None = None, timeline: dict | None = None) -> bytes:
    doc = _setup_document()
    _heading(doc, title, size=16, space_after=4 if project else 10)
    if project:
        period = (f"   |   기간: {project['start_date']} ~ {project['end_date']}"
                  if project.get("start_date") and project.get("end_date") else "")
        _heading(doc, f"프로젝트명: {project.get('project_name', '')}{period}", size=12, space_after=10)
        if project.get("items"):
            _add_profile_section(doc, project)
    if timeline:
        _add_timeline_section(doc, timeline, categories)

    for i, category in enumerate(c for c in categories if c["items"]):
        if i:
            doc.add_paragraph()  # 표 사이 간격 (마지막 뒤에는 넣지 않아 빈 페이지 방지)
        _heading(doc, category["title"], size=11.5, space_after=4)
        _add_table(doc, category["items"])

    _add_basis_appendix(doc, categories)
    _add_forms_appendix(doc, categories)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
